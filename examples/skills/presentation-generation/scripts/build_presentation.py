#!/usr/bin/env python3
"""Build a polished, editable PPTX from a bounded JSON brief."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MAX_INPUT_BYTES = 1_048_576
MAX_SLIDES = 50
MAX_TITLE_CHARS = 80
MAX_BODY_CHARS = 240
HEX_COLOR = re.compile(r"^[0-9A-Fa-f]{6}$")
LAYOUTS = {"title", "section", "bullets", "two_column", "metrics"}

DEFAULT_THEME = {
    "primary": "155EEF",
    "accent": "12B76A",
    "background": "F8FAFC",
    "surface": "FFFFFF",
    "text": "101828",
    "muted": "667085",
}


class BriefError(ValueError):
    """The input brief cannot be represented safely by this builder."""


def color(value: str) -> RGBColor:
    if not isinstance(value, str) or not HEX_COLOR.fullmatch(value):
        raise BriefError(f"invalid RGB color: {value!r}")
    return RGBColor.from_string(value.upper())


def text(value: Any, field: str, maximum: int = MAX_BODY_CHARS) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise BriefError(f"{field} must be a string")
    value = value.strip()
    if len(value) > maximum:
        raise BriefError(f"{field} exceeds {maximum} characters")
    return value


def required_text(value: Any, field: str, maximum: int = MAX_BODY_CHARS) -> str:
    value = text(value, field, maximum)
    if not value:
        raise BriefError(f"{field} is required")
    return value


def string_list(value: Any, field: str, maximum_items: int = 8) -> list[str]:
    if value is None:
        return []
    if not isinstance(value, list) or len(value) > maximum_items:
        raise BriefError(f"{field} must be a list with at most {maximum_items} items")
    return [required_text(item, f"{field}[{index}]") for index, item in enumerate(value)]


def load_brief(path: Path) -> dict[str, Any]:
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise BriefError(f"input exceeds {MAX_INPUT_BYTES} bytes")
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise BriefError(f"cannot read JSON brief: {exc}") from exc
    if not isinstance(value, dict):
        raise BriefError("brief root must be an object")
    required_text(value.get("title"), "title", MAX_TITLE_CHARS)
    slides = value.get("slides")
    if not isinstance(slides, list) or not slides or len(slides) > MAX_SLIDES:
        raise BriefError(f"slides must contain between 1 and {MAX_SLIDES} entries")
    for index, slide in enumerate(slides):
        if not isinstance(slide, dict):
            raise BriefError(f"slides[{index}] must be an object")
        if slide.get("layout") not in LAYOUTS:
            raise BriefError(f"slides[{index}].layout is unsupported")
    return value


def resolve_theme(brief: dict[str, Any]) -> dict[str, RGBColor]:
    overrides = brief.get("theme") or {}
    if not isinstance(overrides, dict):
        raise BriefError("theme must be an object")
    unknown = set(overrides) - set(DEFAULT_THEME)
    if unknown:
        raise BriefError(f"unknown theme keys: {', '.join(sorted(unknown))}")
    return {name: color(str(overrides.get(name, value))) for name, value in DEFAULT_THEME.items()}


def fill_shape(shape: Any, value: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = value
    shape.line.fill.background()


def add_text(
    slide: Any,
    value: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: int,
    rgb: RGBColor,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font: str = "Aptos",
) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb
    return box


def add_title(slide: Any, value: str, theme: dict[str, RGBColor]) -> None:
    add_text(slide, required_text(value, "slide.title", MAX_TITLE_CHARS), 0.75, 0.55, 11.8, 0.65,
             size=26, rgb=theme["text"], bold=True)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.28), Inches(0.7), Inches(0.06))
    fill_shape(accent, theme["accent"])


def add_footer(slide: Any, footer: str, number: int, theme: dict[str, RGBColor]) -> None:
    add_text(slide, footer, 0.75, 7.05, 10.8, 0.22, size=8, rgb=theme["muted"])
    add_text(slide, f"{number:02d}", 11.9, 7.02, 0.65, 0.25, size=8,
             rgb=theme["muted"], align=PP_ALIGN.RIGHT)


def add_bullet_list(slide: Any, bullets: list[str], x: float, y: float, width: float, height: float,
                    theme: dict[str, RGBColor], size: int = 19) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.04)
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = theme["text"]
        paragraph.space_after = Pt(14)
        paragraph.line_spacing = 1.12


def render_title(slide: Any, spec: dict[str, Any], theme: dict[str, RGBColor]) -> None:
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SLIDE_HEIGHT)
    fill_shape(band, theme["accent"])
    add_text(slide, text(spec.get("eyebrow") or "PRESENTATION", "eyebrow"), 0.9, 1.2, 5.5, 0.35,
             size=11, rgb=theme["primary"], bold=True)
    add_text(slide, required_text(spec.get("title"), "slide.title", MAX_TITLE_CHARS), 0.9, 1.75, 11.2, 1.7,
             size=38, rgb=theme["text"], bold=True)
    add_text(slide, text(spec.get("subtitle"), "slide.subtitle"), 0.95, 3.7, 10.5, 0.9,
             size=19, rgb=theme["muted"])
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(5.4), Inches(2.0), Inches(0.08))
    fill_shape(rule, theme["primary"])


def render_section(slide: Any, spec: dict[str, Any], theme: dict[str, RGBColor]) -> None:
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(0.75), Inches(11.9), Inches(5.9))
    fill_shape(panel, theme["primary"])
    add_text(slide, text(spec.get("eyebrow") or "SECTION", "eyebrow"), 1.25, 1.25, 5.0, 0.35,
             size=11, rgb=theme["surface"], bold=True)
    add_text(slide, required_text(spec.get("title"), "slide.title", MAX_TITLE_CHARS), 1.25, 2.0, 10.5, 1.5,
             size=34, rgb=theme["surface"], bold=True)
    add_text(slide, text(spec.get("subtitle"), "slide.subtitle"), 1.28, 4.0, 9.8, 1.1,
             size=18, rgb=theme["surface"])


def render_bullets(slide: Any, spec: dict[str, Any], theme: dict[str, RGBColor]) -> None:
    add_title(slide, spec.get("title"), theme)
    add_text(slide, text(spec.get("subtitle"), "slide.subtitle"), 0.78, 1.5, 11.4, 0.55,
             size=14, rgb=theme["muted"])
    bullets = string_list(spec.get("bullets"), "bullets", 7)
    if not bullets:
        raise BriefError("bullets layout requires at least one bullet")
    add_bullet_list(slide, bullets, 1.0, 2.15, 11.0, 4.4, theme)


def column(spec: Any, field: str) -> tuple[str, list[str]]:
    if not isinstance(spec, dict):
        raise BriefError(f"{field} must be an object")
    title_value = required_text(spec.get("title"), f"{field}.title", MAX_TITLE_CHARS)
    bullets = string_list(spec.get("bullets"), f"{field}.bullets", 6)
    if not title_value or not bullets:
        raise BriefError(f"{field} requires title and bullets")
    return title_value, bullets


def render_two_column(slide: Any, spec: dict[str, Any], theme: dict[str, RGBColor]) -> None:
    add_title(slide, spec.get("title"), theme)
    for index, field in enumerate(("left", "right")):
        heading, bullets = column(spec.get(field), field)
        x = 0.75 + index * 6.05
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.7), Inches(5.65), Inches(4.9))
        fill_shape(card, theme["surface"])
        card.line.color.rgb = theme["primary"] if index else theme["accent"]
        card.line.width = Pt(1.25)
        add_text(slide, heading, x + 0.35, 2.05, 4.95, 0.55, size=20,
                 rgb=theme["primary"] if index else theme["accent"], bold=True)
        add_bullet_list(slide, bullets, x + 0.35, 2.85, 4.9, 3.15, theme, size=16)


def render_metrics(slide: Any, spec: dict[str, Any], theme: dict[str, RGBColor]) -> None:
    add_title(slide, spec.get("title"), theme)
    metrics = spec.get("metrics")
    if not isinstance(metrics, list) or not 1 <= len(metrics) <= 4:
        raise BriefError("metrics must contain between 1 and 4 entries")
    gap = 0.28
    width = (11.8 - gap * (len(metrics) - 1)) / len(metrics)
    for index, metric in enumerate(metrics):
        if not isinstance(metric, dict):
            raise BriefError(f"metrics[{index}] must be an object")
        x = 0.75 + index * (width + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(width), Inches(4.6))
        fill_shape(card, theme["surface"])
        add_text(slide, required_text(metric.get("value"), f"metrics[{index}].value", 32), x + 0.25, 2.15,
                 width - 0.5, 1.05, size=31, rgb=theme["primary"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, required_text(metric.get("label"), f"metrics[{index}].label", 80), x + 0.25, 3.35,
                 width - 0.5, 0.65, size=17, rgb=theme["text"], bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, text(metric.get("detail"), f"metrics[{index}].detail"), x + 0.3, 4.25,
                 width - 0.6, 1.15, size=12, rgb=theme["muted"], align=PP_ALIGN.CENTER)


RENDERERS = {
    "title": render_title,
    "section": render_section,
    "bullets": render_bullets,
    "two_column": render_two_column,
    "metrics": render_metrics,
}


def build(brief: dict[str, Any], output: Path) -> None:
    theme = resolve_theme(brief)
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    presentation.core_properties.title = text(brief.get("title"), "title", MAX_TITLE_CHARS)
    presentation.core_properties.subject = "Editable presentation generated by WeKnora"
    presentation.core_properties.author = text(brief.get("author"), "author", 120)
    footer = text(brief.get("footer"), "footer", 160)

    for number, spec in enumerate(brief["slides"], 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = theme["background"]
        RENDERERS[spec["layout"]](slide, spec, theme)
        add_footer(slide, footer, number, theme)
        notes = text(spec.get("notes"), f"slides[{number - 1}].notes", 2_000)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="UTF-8 JSON brief")
    parser.add_argument("--output", required=True, type=Path, help="destination .pptx")
    args = parser.parse_args()
    if args.output.suffix.lower() != ".pptx":
        parser.error("--output must end with .pptx")
    try:
        brief = load_brief(args.input)
        build(brief, args.output)
    except (BriefError, OSError) as exc:
        parser.error(str(exc))
    print(json.dumps({
        "output": str(args.output),
        "slides": len(brief["slides"]),
        "bytes": args.output.stat().st_size,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
