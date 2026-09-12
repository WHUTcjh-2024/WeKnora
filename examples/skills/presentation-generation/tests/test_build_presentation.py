from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree
from pathlib import Path

from pptx import Presentation

SKILL_ROOT = Path(__file__).resolve().parents[1]
BUILDER = SKILL_ROOT / "scripts" / "build_presentation.py"
EXAMPLE = SKILL_ROOT / "assets" / "example.json"


def load_builder():
    spec = importlib.util.spec_from_file_location("presentation_builder", BUILDER)
    if spec is None or spec.loader is None:
        raise RuntimeError("cannot load presentation builder")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


builder = load_builder()


class PresentationBuilderTest(unittest.TestCase):
    def build_example(self, directory: Path) -> tuple[Path, dict[str, object]]:
        output = directory / "topic2.pptx"
        completed = subprocess.run(
            [sys.executable, str(BUILDER), "--input", str(EXAMPLE), "--output", str(output)],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        return output, json.loads(completed.stdout)

    def test_example_is_valid_editable_unicode_pptx(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, summary = self.build_example(Path(tmp))
            self.assertEqual(summary["slides"], 5)
            self.assertEqual(summary["bytes"], output.stat().st_size)
            self.assertGreater(output.stat().st_size, 10_000)

            deck = Presentation(output)
            self.assertEqual(len(deck.slides), 5)
            for slide in deck.slides:
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                    self.assertLessEqual(shape.top + shape.height, deck.slide_height)
            for slide in deck.slides:
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                    self.assertLessEqual(shape.top + shape.height, deck.slide_height)
            slide_text = [
                shape.text
                for slide in deck.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            ]
            joined = "\n".join(slide_text)
            self.assertIn("WeKnora 犀牛鸟课题二", joined)
            self.assertIn("Safe Files", joined)
            self.assertIn("跨租户与跨会话访问", joined)

            editable = next(shape for shape in deck.slides[0].shapes if shape.has_text_frame and "WeKnora" in shape.text)
            editable.text_frame.paragraphs[0].runs[0].text = "可编辑标题"
            edited = Path(tmp) / "edited.pptx"
            deck.save(edited)
            self.assertIn("可编辑标题", "\n".join(
                shape.text for shape in Presentation(edited).slides[0].shapes if shape.has_text_frame
            ))

    def test_package_contains_slide_parts_and_no_rasterized_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            output, _ = self.build_example(Path(tmp))
            with zipfile.ZipFile(output) as archive:
                names = set(archive.namelist())
                self.assertIn("[Content_Types].xml", names)
                self.assertIn("ppt/presentation.xml", names)
                self.assertEqual(len([
                    name for name in names
                    if name.startswith("ppt/slides/slide") and name.endswith(".xml")
                ]), 5)
                self.assertFalse(any(name.startswith("ppt/media/") for name in names))
                drawing_namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                editable_text = []
                for name in names:
                    if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                        root = ElementTree.fromstring(archive.read(name))
                        editable_text.extend(node.text or "" for node in root.iter(f"{drawing_namespace}t"))
                joined = "\n".join(editable_text)
                self.assertIn("WeKnora", joined)
                self.assertIn("\u7280\u725b\u9e1f", joined)

    def test_validation_rejects_unknown_layout_and_oversized_deck(self):
        valid = json.loads(EXAMPLE.read_text(encoding="utf-8"))
        valid["slides"][0]["layout"] = "arbitrary_code"
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.json"
            path.write_text(json.dumps(valid), encoding="utf-8")
            with self.assertRaisesRegex(builder.BriefError, "unsupported"):
                builder.load_brief(path)

        valid = json.loads(EXAMPLE.read_text(encoding="utf-8"))
        valid["slides"] = valid["slides"] * 11
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "too-many.json"
            path.write_text(json.dumps(valid), encoding="utf-8")
            with self.assertRaisesRegex(builder.BriefError, "between 1 and 50"):
                builder.load_brief(path)


if __name__ == "__main__":
    unittest.main()
